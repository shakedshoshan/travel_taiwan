import openpyxl,json,re,hashlib,collections,os
W='/workspace/'; O=W+'taiwan-data/'
C=[ # id,he,en,lat,lng
('taipei','טייפה','Taipei',25.0478,121.5170),('new-taipei','ניו טייפה','New Taipei',25.0120,121.4650),
('keelung','קילונג','Keelung',25.1283,121.7419),('yilan','יילאן','Yilan',24.7570,121.7530),
('taoyuan','טאיואן','Taoyuan',24.9937,121.3010),('hsinchu','שינג׳ו','Hsinchu',24.8016,120.9716),
('miaoli','מיאולי','Miaoli',24.5602,120.8214),('taichung','טאיצ׳ונג','Taichung',24.1477,120.6736),
('changhua','ג׳אנגהואה','Changhua',24.0809,120.5385),('nantou','נאנטואו','Nantou',23.9157,120.6639),
('sun-moon-lake','אגם הירח והשמש','Sun Moon Lake',23.8650,120.9150),('chiayi','צ׳יאיי','Chiayi',23.4801,120.4491),
('alishan','אלישאן','Alishan',23.5100,120.8040),('tainan','טאינאן','Tainan',22.9908,120.2133),
('kaohsiung','קאושיונג','Kaohsiung',22.6273,120.3014),('pingtung','פינגטונג','Pingtung',22.6690,120.4880),
('hualien','חואליאן','Hualien',23.9872,121.6015),('taitung','טאיטונג','Taitung',22.7583,121.1444)]
def cid(s):
  s=(s or '').strip()
  if any(k in s for k in ['אגם','סאן מון','אידאשאו']): return 'sun-moon-lake'
  if 'אלישאן' in s: return 'alishan'
  if 'לוקאנג' in s: return 'changhua'
  rules=[('ניו טייפה','new-taipei'),('טייפה','taipei'),('טאיפיי','taipei'),('קילונג','keelung'),('יילאן','yilan'),('טאויואן','taoyuan'),('טאיואן','taoyuan'),('שינג׳ו','hsinchu'),('מיאולי','miaoli'),('טאיצ','taichung'),('ג׳אנגהואה','changhua'),('נאנטואו','nantou'),('צ׳יאיי','chiayi'),('טאינאן','tainan'),('קאושיונג','kaohsiung'),('פינגטונג','pingtung'),('חואליאן','hualien'),('הואליין','hualien'),('טאיטונג','taitung')]
  for k,v in rules:
    if s.startswith(k): return v
  return None
unm=collections.Counter()
def zh(s):
  m=re.findall(r'[\u4e00-\u9fff]+',s or ''); return m[0] if m else None
def slug(p,name,i): return f"{p}-{i:03d}-"+hashlib.md5(name.encode()).hexdigest()[:6]
def price(t):
  t=str(t or ''); nums=[int(x.replace(',','')) for x in re.findall(r'\d[\d,]*',t.split('|')[0]) if int(x.replace(',',''))>=10]
  return (min(nums) if nums else None,max(nums) if nums else None)
def ver(*ts): return not any(('לא אומת' in str(t or '') or str(t or '').strip().startswith('~')) for t in ts)
def rows(f):
  ws=openpyxl.load_workbook(W+f).active; return ws,[(i,r) for i,r in enumerate(ws.iter_rows(min_row=2,values_only=True),2) if r[0]]
def city(c,o):
  o['cityId']=cid(c)
  if not o['cityId']: unm[c]+=1
  cc=next((x for x in C if x[0]==o['cityId']),None)
  o.update(lat=cc[3] if cc else None,lng=cc[4] if cc else None,approxLocation=True)
ws,R=rows('taiwan-attractions.xlsx'); A=[]
for n,(i,r) in enumerate(R):
  o=dict(id=slug('attr',r[0],n),name=r[0],nameZh=zh(r[0]),city=r[1],description=r[2],importance=r[3],isNature=ws.cell(i,1).fill.fgColor.rgb=='00C6EFCE'); city(r[1],o); A.append(o)
ws,R=rows('taiwan-food.xlsx'); TM={'שוק':'market','זול':'cheap','מישלן ביב':'michelin-bib','מקומי נסתר':'hidden-local'}; F=[]
for n,(i,r) in enumerate(R):
  mn,mx=price(r[6]); o=dict(id=slug('food',r[0],n),name=r[0],nameZh=zh(r[0]),city=r[1],typeHe=r[2],type=TM.get(r[2],'other'),description=r[3],hours=r[4],importance=r[5],priceText=r[6],priceMin=mn,priceMax=mx,priceUnit='NT$ per person',hoursVerified='לבדוק' not in str(r[4] or ''),priceVerified=ver(r[6]),verified=('לבדוק' not in str(r[4] or '')) and ver(r[6])); city(r[1],o); F.append(o)
ws,R=rows('taiwan-hotels.xlsx'); HT={'פינוק':'splurge','זול':'cheap','איכותי':'quality'}; H=[]
R=[x for x in R if x[1][2]!='—']
for n,(i,r) in enumerate(R):
  mn,mx=price(r[3]); t=next((v for k,v in HT.items() if k in (r[2] or '')),'other')
  o=dict(id=slug('hotel',r[0],n),name=r[0],nameZh=zh(r[0]),city=r[1],tierHe=r[2],tier=t,priceText=r[3],priceMin=mn,priceMax=mx,priceUnit='NT$ per night',description=r[4],source=r[5],verified=ver(r[3])); city(r[1],o); H.append(o)
wb=openpyxl.load_workbook(W+'taiwan-route.xlsx'); RT=[]
RC=['taipei','hualien','taitung','kaohsiung','tainan','alishan','sun-moon-lake','taipei','taichung','taipei']
for n,r in enumerate(wb['מסלול'].iter_rows(min_row=2,values_only=True)):
  RT.append(dict(order=n+1,id=f'leg-{n+1:02d}',focus=r[0],cityId=RC[n],days=r[1],dates=r[2],anchor=r[3],notes=r[4]))
audit=[dict(topic=r[0],status=r[1],finding=r[2]) for r in wb['בקרה'].iter_rows(min_row=2,values_only=True) if r[0]]
cities=[]
for c in C:
  legs=[l for l in RT if l['cityId']==c[0]]
  cities.append(dict(id=c[0],nameHe=c[1],nameEn=c[2],lat=c[3],lng=c[4],onRoute=bool(legs),routeOrder=[l['order'] for l in legs] or None,dates=[l['dates'] for l in legs],days=sum(l['days'] for l in legs)))
D=lambda f,x: json.dump(x,open(O+f,'w'),ensure_ascii=False,indent=2)
D('attractions.json',A);D('food.json',F);D('hotels.json',H);D('route.json',dict(legs=RT,audit=audit));D('cities.json',cities)
ids={c[0] for c in C}; bad=[x['name'] for x in A+F+H if x['cityId'] not in ids]
print(len(A),sum(a['isNature'] for a in A),len(F),collections.Counter(f['type'] for f in F),sum(f['verified'] for f in F),len(H),collections.Counter(h['tier'] for h in H),sum(h['verified'] for h in H),sum(h['priceMin'] is not None for h in H),len(RT),len(audit),'unmapped',unm,bad)
print(collections.Counter(r[2] for _,r in rows('taiwan-food.xlsx')[1]),collections.Counter(r[2] for _,r in rows('taiwan-hotels.xlsx')[1]))
